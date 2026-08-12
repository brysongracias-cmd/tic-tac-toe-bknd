import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# (method, endpoint, description, status_code, pass_fail, reason)
results = [
    ("GET", "/health", "Health check", 200, "PASS", "Real uvicorn boot verified"),
    ("GET", "/", "Root status", 200, "PASS", "Endpoint available"),
    ("POST", "/api/v1/auth/register", "Register a player account", 201, "PASS", "pytest tests/test_auth.py::test_register"),
    ("POST", "/api/v1/auth/login", "Login and receive JWT", 200, "PASS", "pytest tests/test_auth.py::test_login"),
    ("GET", "/api/v1/auth/me", "Read authenticated user", 200, "PASS", "pytest tests/test_auth.py::test_me"),
    ("GET", "/api/v1/auth/me", "Reject invalid bearer token", 401, "PASS", "pytest tests/test_auth.py::test_invalid_token"),
    ("POST", "/api/v1/games/", "Create tic-tac-toe game", 201, "PASS", "pytest tests/test_games.py::test_create_game"),
    ("GET", "/api/v1/games/", "List authenticated user's games", 200, "PASS", "pytest tests/test_games.py::test_list_games"),
    ("GET", "/api/v1/games/{game_id}", "Get a game", 200, "PASS", "pytest tests/test_games.py::test_get_game"),
    ("PATCH", "/api/v1/games/{game_id}", "Update a game", 200, "PASS", "pytest tests/test_games.py::test_update_game"),
    ("DELETE", "/api/v1/games/{game_id}", "Delete a game", 204, "PASS", "pytest tests/test_games.py::test_delete_game"),
    ("GET", "/api/v1/games/{game_id}", "Deleted game returns not found", 404, "PASS", "pytest tests/test_games.py::test_delete_game"),
    ("POST", "/api/v1/games/{game_id}/moves", "Make next tic-tac-toe move", 201, "PASS", "pytest tests/test_games.py::test_make_move"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "API Test Report"
hf = Font(bold=True, color="FFFFFF", size=11)
hbg = PatternFill("solid", fgColor="2F5496")
pg = PatternFill("solid", fgColor="C6EFCE")
fr = PatternFill("solid", fgColor="FFC7CE")
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
lft = Alignment(horizontal="left", vertical="center", wrap_text=True)
t = Side(style="thin")
bdr = Border(left=t, right=t, top=t, bottom=t)
for c, h in enumerate(["#", "Method", "Endpoint", "Description", "Status Code", "Pass/Fail", "Reason"], 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = hf; cell.fill = hbg; cell.alignment = ctr; cell.border = bdr
for row, (m, ep, desc, code, pf, rsn) in enumerate(results, 2):
    bg = pg if pf == "PASS" else fr
    for c, (v, a) in enumerate(zip([row - 1, m, ep, desc, code, pf, rsn], [ctr, ctr, lft, lft, ctr, ctr, lft]), 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.fill = bg; cell.alignment = a; cell.border = bdr
        if c == 6:
            cell.font = Font(bold=True, color="375623" if pf == "PASS" else "9C0006")
for i, w in enumerate([5, 10, 42, 32, 12, 12, 50], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
wb.save("api_test_report.xlsx")
print("Saved: api_test_report.xlsx")
